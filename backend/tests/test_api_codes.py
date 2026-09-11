"""API tests for the codes router: CRUD, frequency, subtree delete, merge, docx export,
AI suggestions, and the ai_suggested flag.

Note the trailing-slash paths: the codes router is mounted at
/projects/{id}/codes/ with the list and create routes registered as "/".
"""
import io

import docx as python_docx

import app.models as models
from app.services.llm_service import llm_service, LLMResponseError
from tests.helpers import make_code, make_document, make_memo, make_segment


def _create_code(client, project_id, name, **extra):
    response = client.post(f"/projects/{project_id}/codes/", json={"name": name, **extra})
    assert response.status_code == 201
    return response.json()


def test_list_codes_reports_segment_frequency(client, db, project, code, document):
    make_segment(db, document, code, 0, 5, "alpha")
    make_segment(db, document, code, 6, 10, "beta")
    make_code(db, project, name="Empty")

    response = client.get(f"/projects/{project.id}/codes/")

    assert response.status_code == 200
    frequencies = {c["name"]: c["frequency"] for c in response.json()}
    assert frequencies == {"Theme": 2, "Empty": 0}


def test_create_code_root_and_subcode(client, project):
    parent = _create_code(client, project.id, "Parent", color="#ABCDEF")

    assert parent["parent_id"] is None
    assert parent["project_id"] == project.id
    assert parent["frequency"] == 0

    child = _create_code(client, project.id, "Child", parent_id=parent["id"])
    assert child["parent_id"] == parent["id"]


def test_update_code_and_404(client, project):
    code = _create_code(client, project.id, "Old", color="#111111")

    response = client.put(f"/projects/{project.id}/codes/{code['id']}",
                          json={"name": "New", "color": "#222222"})
    assert response.status_code == 200
    assert response.json()["name"] == "New"
    assert response.json()["color"] == "#222222"

    assert client.put(f"/projects/{project.id}/codes/9999",
                      json={"name": "X"}).status_code == 404


def test_reorder_codes_updates_parent_and_order(client, project):
    parent = _create_code(client, project.id, "Parent")
    child = _create_code(client, project.id, "Child")

    response = client.put(f"/projects/{project.id}/codes/reorder", json={
        "codes": [{"id": child["id"], "parent_id": parent["id"], "order_index": 5}],
    })

    assert response.status_code == 200
    by_name = {c["name"]: c for c in response.json()}
    assert by_name["Child"]["parent_id"] == parent["id"]
    assert by_name["Child"]["order_index"] == 5


def test_delete_code_cleans_subtree_end_to_end(client, db, project, document):
    """End-to-end version of the repo-level cascade test: deleting a parent
    code through the API must remove its subcodes, their segments, and their
    code memos (SQLite FK cascades never fire, so the repo deletes by hand)."""
    parent = _create_code(client, project.id, "Parent")
    child = _create_code(client, project.id, "Child", parent_id=parent["id"])
    unrelated = _create_code(client, project.id, "Unrelated")

    seg_child = client.post(f"/projects/{project.id}/segments", json={
        "document_id": document.id, "code_id": child["id"],
        "start_char": 0, "end_char": 5, "content": "alpha",
    }).json()
    client.post(f"/projects/{project.id}/segments", json={
        "document_id": document.id, "code_id": unrelated["id"],
        "start_char": 6, "end_char": 10, "content": "beta",
    })
    client.post("/memos", json={
        "text": "child memo", "target_type": "code", "target_id": child["id"],
    })
    keep_memo = client.post("/memos", json={
        "text": "keep", "target_type": "code", "target_id": unrelated["id"],
    }).json()

    response = client.delete(f"/projects/{project.id}/codes/{parent['id']}")

    assert response.status_code == 204
    assert [c["name"] for c in client.get(f"/projects/{project.id}/codes/").json()] == ["Unrelated"]

    remaining = client.get(f"/projects/{project.id}/segments").json()
    assert [s["id"] for s in remaining] != [seg_child["id"]]  # child's segment is gone
    assert len(remaining) == 1

    memos = client.get(f"/projects/{project.id}/memos").json()
    assert [m["id"] for m in memos] == [keep_memo["id"]]


def test_delete_code_404(client, project):
    assert client.delete(f"/projects/{project.id}/codes/9999").status_code == 404


def test_merge_codes_via_api(client, project, document):
    source = _create_code(client, project.id, "Source")
    target = _create_code(client, project.id, "Target")
    seg = client.post(f"/projects/{project.id}/segments", json={
        "document_id": document.id, "code_id": source["id"],
        "start_char": 0, "end_char": 5, "content": "alpha",
    }).json()

    response = client.post(f"/projects/{project.id}/codes/merge", json={
        "source_code_id": source["id"], "target_code_id": target["id"],
        "new_name": "Merged", "new_color": "#00FF00",
    })

    assert response.status_code == 200
    body = response.json()
    assert body["id"] == target["id"]
    assert body["name"] == "Merged"
    assert body["color"] == "#00FF00"

    names = [c["name"] for c in client.get(f"/projects/{project.id}/codes/").json()]
    assert "Source" not in names

    segs = client.get(f"/projects/{project.id}/segments").json()
    assert [s["id"] for s in segs] == [seg["id"]]
    assert segs[0]["code_id"] == target["id"]


def test_merge_missing_code_400(client, project):
    target = _create_code(client, project.id, "Target")
    response = client.post(f"/projects/{project.id}/codes/merge", json={
        "source_code_id": 9999, "target_code_id": target["id"],
    })
    assert response.status_code == 400
    assert response.json()["detail"] == "Merge failed. Ensure both codes exist."


def test_export_codebook_docx(client, db, project):
    parent = make_code(db, project, name="Economy")
    make_code(db, project, name="Jobs", parent=parent)
    make_memo(db, "code", parent.id, text="important memo")

    response = client.get(f"/projects/{project.id}/codes/export/docx")

    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert "Codebook_Test%20Project.docx" in response.headers["content-disposition"]

    document = python_docx.Document(io.BytesIO(response.content))
    texts = [p.text for p in document.paragraphs]

    assert texts[0] == "Codebook: Test Project"
    # code names are indented 4 spaces per hierarchy depth
    assert "Economy" in texts
    assert "    Jobs" in texts
    # the bold "Memo: " prefix run and the memo text form one paragraph
    memo_paragraph = next(p for p in document.paragraphs if p.text == "Memo: important memo")
    assert memo_paragraph.runs[0].text == "Memo: "
    assert memo_paragraph.runs[0].bold


def test_export_codebook_docx_404(client):
    assert client.get("/projects/9999/codes/export/docx").status_code == 404


# --- ai_suggested flag -------------------------------------------------------

def test_create_code_defaults_not_ai_suggested(client, project):
    body = _create_code(client, project.id, "Human")
    assert body["ai_suggested"] is False


def test_create_code_with_ai_suggested_flag(client, db, project):
    body = _create_code(client, project.id, "AI Code", ai_suggested=True)
    assert body["ai_suggested"] is True

    db.expire_all()  # stale-session rule: repos commit on the request session
    persisted = db.query(models.Code).filter(models.Code.id == body["id"]).first()
    assert persisted.ai_suggested is True


def test_update_code_can_unflag_ai_suggested(client, project):
    code = _create_code(client, project.id, "AI Code", ai_suggested=True)

    response = client.put(f"/projects/{project.id}/codes/{code['id']}",
                          json={"ai_suggested": False})
    assert response.status_code == 200
    assert response.json()["ai_suggested"] is False


def test_update_code_omitting_flag_leaves_it_true(client, project):
    code = _create_code(client, project.id, "AI Code", ai_suggested=True)

    response = client.put(f"/projects/{project.id}/codes/{code['id']}",
                          json={"name": "Renamed"})
    assert response.status_code == 200
    assert response.json()["ai_suggested"] is True


# --- AI suggestions endpoint -------------------------------------------------

def _seed_llm_settings(db, api_url="https://api.test/v1", model="test-model",
                        api_key="sk-test"):
    db.add(models.LLMSettings(id=1, api_url=api_url, model_name=model,
                              api_key=api_key))
    db.commit()


def test_suggest_codes_404_unknown_project(client):
    assert client.post("/projects/9999/codes/suggest",
                       json={"excerpt": "text"}).status_code == 404


def test_suggest_codes_409_when_not_configured(client, project):
    response = client.post(f"/projects/{project.id}/codes/suggest",
                           json={"excerpt": "text"})
    assert response.status_code == 409
    assert response.json()["detail"] == "LLM settings are not configured."


def test_suggest_codes_returns_suggestions(client, db, project, monkeypatch):
    _seed_llm_settings(db)
    existing = make_code(db, project, name="Theme")

    monkeypatch.setattr(llm_service, "_chat_completion",
                        lambda *a, **k: '{"suggestions": [{"name": "Theme"},'
                                       ' {"name": "Fresh", "rationale": "new"}]}')

    response = client.post(f"/projects/{project.id}/codes/suggest",
                           json={"excerpt": "some selection"})

    assert response.status_code == 200
    suggestions = response.json()
    assert suggestions[0] == {"name": "Theme", "rationale": None,
                              "existing_code_id": existing.id}
    assert suggestions[1]["existing_code_id"] is None


def test_suggest_codes_502_on_llm_failure(client, db, project, monkeypatch):
    _seed_llm_settings(db)

    def raise_response(*args, **kwargs):
        raise LLMResponseError("LLM did not return valid JSON.")
    monkeypatch.setattr(llm_service, "_chat_completion", raise_response)

    response = client.post(f"/projects/{project.id}/codes/suggest",
                           json={"excerpt": "text"})
    assert response.status_code == 502
    assert response.json()["detail"] == "LLM did not return valid JSON."