"""API tests for the projects router: CRUD, filesystem behavior, search, Excel export."""
import io

import openpyxl

from tests.helpers import make_code, make_document, make_project, make_segment


def test_root_endpoint(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "message" in response.json()


def test_list_projects_empty(client):
    response = client.get("/projects")
    assert response.status_code == 200
    assert response.json() == []


def test_create_project_without_local_path(client):
    response = client.post("/projects", json={
        "name": "Workspace A",
        "description": "first project",
    })
    assert response.status_code == 200
    body = response.json()
    assert body["name"] == "Workspace A"
    assert body["description"] == "first project"
    assert body["local_path"] is None
    # ProjectRepository.create passes last_accessed=None explicitly, but the
    # ORM omits None-valued attributes from the INSERT, so the column default
    # fires and a timestamp lands anyway.
    assert body["last_accessed"] is not None
    assert body["document_count"] == 0
    assert body["code_count"] == 0


def test_create_project_with_local_path_creates_directory(client, tmp_path):
    base = tmp_path / "base"
    response = client.post("/projects", json={
        "name": "My WS",
        "local_path": str(base),
    })
    assert response.status_code == 200
    assert response.json()["local_path"] == str(base / "My WS")
    assert (base / "My WS").is_dir()


def test_create_project_duplicate_base_path_400(client, db, tmp_path):
    # Mimics an imported project whose stored local_path equals the base dir.
    make_project(db, name="Existing", local_path=str(tmp_path))
    response = client.post("/projects", json={"name": "Other", "local_path": str(tmp_path)})
    assert response.status_code == 400
    assert "already using this folder" in response.json()["detail"]


def test_create_project_existing_target_dir_400(client, tmp_path):
    (tmp_path / "Taken").mkdir()
    response = client.post("/projects", json={"name": "Taken", "local_path": str(tmp_path)})
    assert response.status_code == 400
    assert "already exists" in response.json()["detail"]


def test_get_project_returns_name_description_and_default_templates(client):
    from app.llm_defaults import DEFAULT_SYSTEM_PROMPT, DEFAULT_USER_PROMPT
    project_id = client.post("/projects", json={"name": "Workspace A"}).json()["id"]
    response = client.get(f"/projects/{project_id}")
    assert response.status_code == 200
    assert response.json() == {
        "name": "Workspace A",
        "description": None,
        "llm_system_prompt": DEFAULT_SYSTEM_PROMPT,
        "llm_user_prompt": DEFAULT_USER_PROMPT,
    }


def test_update_project_persists_llm_templates(client, db):
    project_id = client.post("/projects", json={"name": "T"}).json()["id"]

    response = client.put(f"/projects/{project_id}", json={
        "name": "T",
        "llm_system_prompt": "custom system {excerpt}",
        "llm_user_prompt": "custom user {codes} {excerpt}",
    })
    assert response.status_code == 200

    body = client.get(f"/projects/{project_id}").json()
    assert body["llm_system_prompt"] == "custom system {excerpt}"
    assert body["llm_user_prompt"] == "custom user {codes} {excerpt}"

    # PUT without the template fields leaves them unchanged (non-None copy)
    client.put(f"/projects/{project_id}", json={"name": "T2"})
    body = client.get(f"/projects/{project_id}").json()
    assert body["llm_system_prompt"] == "custom system {excerpt}"


def test_get_project_404(client):
    response = client.get("/projects/9999")
    assert response.status_code == 404
    assert response.json()["detail"] == "Project not found"


def test_update_project(client):
    project_id = client.post("/projects", json={"name": "Old"}).json()["id"]

    response = client.put(f"/projects/{project_id}", json={"name": "New", "description": "d"})
    assert response.status_code == 200
    assert response.json()["name"] == "New"

    assert client.get(f"/projects/{project_id}").json()["name"] == "New"
    assert client.put("/projects/9999", json={"name": "X"}).status_code == 404


def test_delete_project_removes_directory(client, tmp_path):
    project_id = client.post("/projects", json={
        "name": "WS", "local_path": str(tmp_path),
    }).json()["id"]
    ws_dir = tmp_path / "WS"
    (ws_dir / "note.txt").write_text("data")

    response = client.delete(f"/projects/{project_id}")

    assert response.status_code == 200
    assert not ws_dir.exists()      # shutil.rmtree'd the joined local_path
    assert tmp_path.exists()        # only the project subfolder is removed
    assert client.get(f"/projects/{project_id}").status_code == 404


def test_delete_project_404(client):
    assert client.delete("/projects/9999").status_code == 404


# --- search -----------------------------------------------------------------

def test_search_returns_all_occurrences_with_context(client, db, project):
    make_document(db, project, filename="doc.txt", content="alpha beta alpha")

    response = client.get(f"/projects/{project.id}/search", params={"query": "alpha"})

    assert response.status_code == 200
    results = response.json()
    assert len(results) == 2
    first, second = results
    assert [first["start_char"], second["start_char"]] == [0, 11]
    assert first["end_char"] == 5
    assert first["document_filename"] == "doc.txt"
    assert first["context"] == "alpha beta alpha"   # whole doc fits in ±100 chars
    assert first["match_offset"] == 0
    assert second["match_offset"] == 11
    assert first["position_label"] == "Char 0"
    assert first["query_length"] == 5
    assert first["is_pdf"] is False


def test_search_pdf_documents_report_page_numbers(client, db, project):
    content = "x" * 2100 + "needle"
    make_document(db, project, filename="scan.pdf", content=content, type="pdf")

    results = client.get(f"/projects/{project.id}/search", params={"query": "needle"}).json()

    assert len(results) == 1
    assert results[0]["position_label"] == "Page 2"
    assert results[0]["is_pdf"] is True
    assert results[0]["context"].startswith("...")    # ellipsis shift reflected in offset
    assert results[0]["match_offset"] == 103


def test_search_blank_query_returns_empty_list(client, project):
    response = client.get(f"/projects/{project.id}/search", params={"query": ""})
    assert response.status_code == 200
    assert response.json() == []


def test_search_missing_query_param_is_422(client, project):
    assert client.get(f"/projects/{project.id}/search").status_code == 422


def test_search_project_404(client):
    assert client.get("/projects/9999/search", params={"query": "x"}).status_code == 404


# --- Excel export -----------------------------------------------------------

def test_export_excel_returns_valid_xlsx(client, db, project):
    doc = make_document(db, project, filename="interview.txt", content="some text here")
    parent = make_code(db, project, name="Economy")
    child = make_code(db, project, name="Jobs", parent=parent)
    make_segment(db, doc, parent, 0, 4, "some")
    make_segment(db, doc, child, 5, 9, "text")

    response = client.get(f"/projects/{project.id}/export/excel")

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "Test%20Project_Statistics.xlsx" in response.headers["content-disposition"]

    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet.title == "Interview Statistics"
    assert [c.value for c in sheet[1]] == [
        "Document Name", "Code Name", "Parent Code", "The Text Segment", "Timestamp",
    ]

    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    by_code = {row[1]: row for row in rows}
    assert by_code["Jobs"][2] == "Economy"     # parent code resolved
    assert by_code["Economy"][2] == "N/A"      # root code has no parent
    assert by_code["Jobs"][3] == "text"
    assert by_code["Jobs"][0] == "interview.txt"


def test_export_excel_respects_docs_and_codes_filters(client, db, project):
    doc = make_document(db, project, filename="interview.txt", content="some text here")
    parent = make_code(db, project, name="Economy")
    child = make_code(db, project, name="Jobs", parent=parent)
    make_segment(db, doc, parent, 0, 4, "some")
    make_segment(db, doc, child, 5, 9, "text")

    response = client.get(f"/projects/{project.id}/export/excel",
                           params={"codes": str(child.id)})
    rows = list(openpyxl.load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=2, values_only=True))
    assert [row[1] for row in rows] == ["Jobs"]

    # non-numeric items are ignored, so the filter is simply not applied
    response = client.get(f"/projects/{project.id}/export/excel",
                          params={"docs": "abc"})
    rows = list(openpyxl.load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2


def test_export_excel_project_404(client):
    assert client.get("/projects/9999/export/excel").status_code == 404