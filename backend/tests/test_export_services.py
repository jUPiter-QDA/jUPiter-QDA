"""Tests for the export/import services and the REFI-QDPX endpoints.

xlsx/docx units drive the services directly (the API surface is already
covered in test_api_projects/test_api_codes); REFI export/import go through
the endpoints because that is where the zip plumbing lives."""
import io
import zipfile
import xml.etree.ElementTree as ET

import docx as python_docx
import openpyxl
import pytest

import app.models as models
from app.services.csv_service import create_codebook_csv
from app.services.docx_service import create_codebook_docx
from app.services.xlsx_service import create_codebook_xlsx
from tests.helpers import build_qdpx_zip, make_code, make_document, make_folder, make_memo, make_project, make_segment


# --- csv ---------------------------------------------------------------------

def test_create_codebook_csv_is_broken(db, project):
    """KNOWN BUG: create_codebook_csv reads project.segments, which does not
    exist on the Project model (segments belong to documents and codes), so
    the csv export always raises. The endpoint test in test_api_segments
    xfails strictly on this; fix the service and both flip."""
    with pytest.raises(AttributeError):
        create_codebook_csv(project)


# --- xlsx --------------------------------------------------------------------

def test_create_codebook_xlsx_returns_stream_and_quoted_filename(db, project, code, document):
    segment = make_segment(db, document, code)
    db.expire(project)

    output, safe_filename = create_codebook_xlsx(project, [code], [segment])

    assert safe_filename == "Test%20Project_Statistics.xlsx"
    sheet = openpyxl.load_workbook(io.BytesIO(output.getvalue())).active
    assert sheet.title == "Interview Statistics"
    assert sheet.freeze_panes == "A2"


# --- docx --------------------------------------------------------------------

def test_create_codebook_docx_builds_hierarchy(db, project, code, document):
    child = make_code(db, project, name="Jobs", parent=code)
    make_memo(db, "code", code.id, text="root code memo")
    make_memo(db, "code", child.id, text="child code memo")
    db.expire(project)

    stream, safe_filename = create_codebook_docx(db, project)

    assert safe_filename == "Codebook_Test%20Project.docx"
    document_out = python_docx.Document(io.BytesIO(stream.getvalue()))
    texts = [p.text for p in document_out.paragraphs]
    date_line = f"Exported from jUPiter QDA on {project.last_accessed.strftime('%B %d, %Y')}"

    assert texts[0] == "Codebook: Test Project"
    assert date_line in texts
    assert "Theme" in texts                      # root at depth 0
    assert "    Jobs" in texts                   # child indented 4 spaces
    assert "Memo: root code memo" in texts
    assert "Memo: child code memo" in texts


def test_create_codebook_docx_crashes_on_null_last_accessed(db, project):
    """LATENT BUG: docx export calls project.last_accessed.strftime(...)
    unguarded, so a NULL last_accessed crashes it. No API path currently
    produces that state (the ORM omits None from INSERTs, so the column
    default fires), but an explicit NULL update would trip this."""
    project.last_accessed = None
    db.commit()

    with pytest.raises(AttributeError):
        create_codebook_docx(db, project)


# --- REFI export -------------------------------------------------------------

def test_export_refi_returns_valid_qdpx_package(client, db, project, code, document):
    child = make_code(db, project, name="Jobs", parent=code)
    segment = make_segment(db, document, code)
    make_memo(db, "project", project.id, text="project note")
    make_memo(db, "code", code.id, text="code note")
    make_memo(db, "segment", segment.id, text="segment note")
    folder = make_folder(db, project, name="Interviews")
    make_document(db, project, filename="in-folder.txt", content="file two", folder=folder)
    db.expire(project)

    response = client.get(f"/projects/{project.id}/export/refi")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert response.headers["content-disposition"] == 'attachment; filename="Test_Project.qdpx"'

    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        names = zf.namelist()
        assert "project.qde" in names
        assert any(n.startswith("Sources/") and n.endswith(".txt") for n in names)

        root = ET.fromstring(zf.read("project.qde"))
    # namespace-stripped lookups
    for elem in root.iter():
        if "}" in elem.tag:
            elem.tag = elem.tag.split("}", 1)[1]

    assert root.tag == "Project"
    assert root.attrib["name"] == "Test Project"
    assert root.find("Description").text == "seed"
    code_names = [c.attrib["name"] for c in root.findall(".//Code")]
    assert code_names == ["Theme", "Jobs"]              # parent before child
    sources = root.findall(".//TextSource")
    assert [s.attrib["name"] for s in sources] == ["doc.txt", "in-folder.txt"]
    selection = sources[0].find("PlainTextSelection")
    assert (selection.attrib["startPosition"], selection.attrib["endPosition"]) == ("0", "5")
    notes = root.findall(".//Note")
    assert len(notes) == 3                             # project + code + segment memos
    set_elem = root.find(".//Set")
    assert set_elem.attrib["name"] == "Interviews"
    assert set_elem.find("MemberSource") is not None   # folder member link


def test_export_refi_404(client):
    assert client.get("/projects/9999/export/refi").status_code == 404


# --- REFI import -------------------------------------------------------------

def test_import_refi_creates_full_project(client, db):
    response = client.post("/projects/import/refi", files={
        "file": ("test.qdpx", build_qdpx_zip(), "application/zip"),
    })

    assert response.status_code == 200
    body = response.json()
    project_id = body["id"]
    assert body["name"] == "Imported Project"
    assert body["description"] == "imported description"
    assert body["document_count"] == 1
    assert body["code_count"] == 2

    documents = client.get(f"/projects/{project_id}/documents/").json()
    assert [d["filename"] for d in documents] == ["interview.txt"]
    fetched = client.get(f"/projects/{project_id}/documents/{documents[0]['id']}").json()
    assert fetched["content"] == "alpha beta gamma"    # read from Sources/doc-1.txt in the zip

    codes = client.get(f"/projects/{project_id}/codes/").json()
    by_name = {c["name"]: c for c in codes}
    assert by_name["Jobs"]["parent_id"] == by_name["Economy"]["id"]

    segments = client.get(f"/projects/{project_id}/segments").json()
    assert len(segments) == 1
    assert segments[0]["content"] == "alpha"           # sliced from the document text
    assert segments[0]["code_id"] == by_name["Economy"]["id"]

    memos = client.get(f"/projects/{project_id}/memos").json()
    by_type = {m["target_type"]: m for m in memos}
    assert by_type["project"]["text"] == "project level memo"
    assert by_type["project"]["target_name"] == "Imported Project"
    assert by_type["code"]["target_name"] == "Economy"
    assert by_type["segment"]["target_name"] == '"alpha"'


def test_import_refi_duplicate_name_gets_suffix(client, db):
    make_project(db, name="Imported Project")

    response = client.post("/projects/import/refi", files={
        "file": ("test.qdpx", build_qdpx_zip(), "application/zip"),
    })

    assert response.status_code == 200
    assert response.json()["name"] == "Imported Project (1)"


def test_import_refi_wrong_extension_400(client):
    response = client.post("/projects/import/refi", files={
        "file": ("package.zip", b"irrelevant", "application/zip"),
    })
    assert response.status_code == 400
    assert response.json()["detail"] == "File must be a .qdpx package"


def test_import_refi_bad_zip_400(client):
    response = client.post("/projects/import/refi", files={
        "file": ("broken.qdpx", b"definitely not a zip", "application/zip"),
    })
    assert response.status_code == 400
    assert response.json()["detail"] == "Invalid QDPX package (Not a valid ZIP file)"


# --- round trip ---------------------------------------------------------------

def test_refi_export_then_import_round_trip(client, db, project, code, document):
    child = make_code(db, project, name="Jobs", parent=code)
    segment = make_segment(db, document, code)
    make_memo(db, "project", project.id, text="project note")
    make_memo(db, "code", code.id, text="code note")
    make_memo(db, "segment", segment.id, text="segment note")
    make_folder(db, project, name="Interviews")
    db.expire(project)

    exported = client.get(f"/projects/{project.id}/export/refi").content

    response = client.post("/projects/import/refi", files={
        "file": ("roundtrip.qdpx", exported, "application/zip"),
    })

    assert response.status_code == 200
    new_id = response.json()["id"]
    assert new_id != project.id
    # the source project still exists in this DB, so the importer's
    # duplicate-name guard renames the copy with a suffix
    assert response.json()["name"] == "Test Project (1)"

    codes = client.get(f"/projects/{new_id}/codes/").json()
    by_name = {c["name"]: c for c in codes}
    assert by_name["Jobs"]["parent_id"] == by_name["Theme"]["id"]

    documents = client.get(f"/projects/{new_id}/documents/").json()
    assert [d["filename"] for d in documents] == ["doc.txt"]
    fetched = client.get(f"/projects/{new_id}/documents/{documents[0]['id']}").json()
    assert fetched["content"] == document.content       # round-tripped through the zip

    segments = client.get(f"/projects/{new_id}/segments").json()
    assert len(segments) == 1
    assert segments[0]["content"] == "alpha"
    assert segments[0]["code_id"] == by_name["Theme"]["id"]

    memos = client.get(f"/projects/{new_id}/memos").json()
    assert {m["target_type"] for m in memos} == {"project", "code", "segment"}
    assert len(memos) == 3

    folders = client.get(f"/projects/{new_id}/folders/").json()
    assert [f["name"] for f in folders] == ["Interviews"]